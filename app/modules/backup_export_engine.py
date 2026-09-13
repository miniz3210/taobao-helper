"""
Backup & Export Engine Module
Handles full system backups, restores, and Excel/CSV exports
"""
import os
import zipfile
import shutil
from datetime import datetime
from pathlib import Path
from typing import List, Dict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import csv
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.order import Order


class BackupExportEngine:
    """Engine for backup, restore, and export operations"""
    
    def __init__(self, data_dir: str = "/app/data"):
        self.data_dir = Path(data_dir)
        self.db_path = self.data_dir / "taobao_helper.db"
        self.snapshots_dir = self.data_dir / "snapshots"
        self.backups_dir = self.data_dir / "backups"
        self.exports_dir = self.data_dir / "exports"
        
        # Ensure directories exist
        self.backups_dir.mkdir(parents=True, exist_ok=True)
        self.exports_dir.mkdir(parents=True, exist_ok=True)
    
    async def create_full_backup(self) -> str:
        """
        Create a full system backup containing database and snapshots
        Returns the path to the backup zip file
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_filename = f"taobao_backup_{timestamp}.zip"
        backup_path = self.backups_dir / backup_filename
        
        with zipfile.ZipFile(backup_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
            # Add database file
            if self.db_path.exists():
                zipf.write(self.db_path, arcname="taobao_helper.db")
            
            # Add all snapshot images
            if self.snapshots_dir.exists():
                for snapshot_file in self.snapshots_dir.glob("*"):
                    if snapshot_file.is_file():
                        arcname = f"snapshots/{snapshot_file.name}"
                        zipf.write(snapshot_file, arcname=arcname)
        
        return str(backup_path)
    
    async def restore_backup(self, backup_zip_path: str) -> Dict[str, any]:
        """
        Restore system from a backup zip file
        Returns statistics about the restore operation
        """
        backup_path = Path(backup_zip_path)
        
        if not backup_path.exists():
            raise FileNotFoundError(f"Backup file not found: {backup_zip_path}")
        
        stats = {
            "database_restored": False,
            "snapshots_restored": 0,
            "errors": []
        }
        
        try:
            with zipfile.ZipFile(backup_path, 'r') as zipf:
                # Extract database
                if "taobao_helper.db" in zipf.namelist():
                    zipf.extract("taobao_helper.db", path=self.data_dir)
                    stats["database_restored"] = True
                
                # Extract snapshots
                for file_info in zipf.namelist():
                    if file_info.startswith("snapshots/") and not file_info.endswith("/"):
                        zipf.extract(file_info, path=self.data_dir)
                        stats["snapshots_restored"] += 1
        
        except Exception as e:
            stats["errors"].append(str(e))
            raise
        
        return stats
    
    async def export_to_excel(self, session: AsyncSession, output_filename: str = None) -> str:
        """
        Export all historical orders to Excel format
        Returns the path to the generated Excel file
        """
        if output_filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"taobao_history_{timestamp}.xlsx"
        
        output_path = self.exports_dir / output_filename
        
        # Fetch all orders
        result = await session.execute(select(Order).order_by(Order.order_date.desc()))
        orders = result.scalars().all()
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Taobao Purchase History"
        
        # Define headers
        headers = [
            "Taobao Order ID",
            "Item Title",
            "AI Short Title",
            "Price (RMB)",
            "Quantity",
            "Seller Name",
            "Seller Express Tracking No.",
            "Carrier",
            "Order Date",
            "Pay Date",
            "Ship Date",
            "Arrival/Receive Date",
            "Current Status",
            "Forwarder Tracking No.",
            "Local Snapshot File Name"
        ]
        
        # Style headers
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Write data rows
        for row_num, order in enumerate(orders, 2):
            snapshot_filename = Path(order.snapshot_local_path).name if order.snapshot_local_path else ""
            
            row_data = [
                order.order_id,
                order.item_title,
                order.ai_clean_title or "",
                order.price,
                order.quantity,
                order.seller_name or "",
                order.seller_express_no or "",
                order.carrier or "",
                order.order_date.strftime("%Y-%m-%d %H:%M:%S") if order.order_date else "",
                order.pay_date.strftime("%Y-%m-%d %H:%M:%S") if order.pay_date else "",
                order.ship_date.strftime("%Y-%m-%d %H:%M:%S") if order.ship_date else "",
                order.receive_date.strftime("%Y-%m-%d %H:%M:%S") if order.receive_date else "",
                order.current_status,
                order.forwarder_tracking_no or "",
                snapshot_filename
            ]
            
            for col_num, value in enumerate(row_data, 1):
                ws.cell(row=row_num, column=col_num, value=value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save workbook
        wb.save(output_path)
        
        return str(output_path)
    
    async def export_to_csv(self, session: AsyncSession, output_filename: str = None) -> str:
        """
        Export all historical orders to CSV format
        Returns the path to the generated CSV file
        """
        if output_filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"taobao_history_{timestamp}.csv"
        
        output_path = self.exports_dir / output_filename
        
        # Fetch all orders
        result = await session.execute(select(Order).order_by(Order.order_date.desc()))
        orders = result.scalars().all()
        
        # Define headers
        headers = [
            "Taobao Order ID",
            "Item Title",
            "AI Short Title",
            "Price (RMB)",
            "Quantity",
            "Seller Name",
            "Seller Express Tracking No.",
            "Carrier",
            "Order Date",
            "Pay Date",
            "Ship Date",
            "Arrival/Receive Date",
            "Current Status",
            "Forwarder Tracking No.",
            "Local Snapshot File Name"
        ]
        
        # Write CSV
        with open(output_path, 'w', newline='', encoding='utf-8-sig') as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(headers)
            
            for order in orders:
                snapshot_filename = Path(order.snapshot_local_path).name if order.snapshot_local_path else ""
                
                row_data = [
                    order.order_id,
                    order.item_title,
                    order.ai_clean_title or "",
                    order.price,
                    order.quantity,
                    order.seller_name or "",
                    order.seller_express_no or "",
                    order.carrier or "",
                    order.order_date.strftime("%Y-%m-%d %H:%M:%S") if order.order_date else "",
                    order.pay_date.strftime("%Y-%m-%d %H:%M:%S") if order.pay_date else "",
                    order.ship_date.strftime("%Y-%m-%d %H:%M:%S") if order.ship_date else "",
                    order.receive_date.strftime("%Y-%m-%d %H:%M:%S") if order.receive_date else "",
                    order.current_status,
                    order.forwarder_tracking_no or "",
                    snapshot_filename
                ]
                
                writer.writerow(row_data)
        
        return str(output_path)
    
    async def export_consolidation_sheet(
        self, 
        session: AsyncSession,
        forwarder_tracking_no: str,
        output_filename: str = None
    ) -> str:
        """
        Export customs declaration sheet for a specific forwarder tracking number
        Returns the path to the generated Excel file
        """
        if output_filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"customs_declaration_{forwarder_tracking_no}_{timestamp}.xlsx"
        
        output_path = self.exports_dir / output_filename
        
        # Fetch orders for this forwarder tracking
        result = await session.execute(
            select(Order).where(Order.forwarder_tracking_no == forwarder_tracking_no)
        )
        orders = result.scalars().all()
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Customs Declaration"
        
        # Headers for customs declaration
        headers = [
            "Item No.",
            "Item Name (CN)",
            "Item Name (EN)",
            "Quantity",
            "Unit Price (RMB)",
            "Total Value (RMB)",
            "Category",
            "Material",
            "Brand",
            "Remarks"
        ]
        
        # Style headers
        header_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Write data rows
        total_value = 0
        for idx, order in enumerate(orders, 1):
            item_total = order.price * order.quantity
            total_value += item_total
            
            row_data = [
                idx,
                order.item_title,
                order.ai_clean_title or "",
                order.quantity,
                order.price,
                item_total,
                "",  # Category - to be filled manually
                "",  # Material - to be filled manually
                "",  # Brand - to be filled manually
                order.seller_express_no or ""
            ]
            
            for col_num, value in enumerate(row_data, 1):
                ws.cell(row=idx + 1, column=col_num, value=value)
        
        # Add summary row
        summary_row = len(orders) + 2
        ws.cell(row=summary_row, column=1, value="Total:")
        ws.cell(row=summary_row, column=1).font = Font(bold=True)
        ws.cell(row=summary_row, column=6, value=total_value)
        ws.cell(row=summary_row, column=6).font = Font(bold=True)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save workbook
        wb.save(output_path)
        
        return str(output_path)
